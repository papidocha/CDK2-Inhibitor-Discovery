"""
Build the Word summary tables (Hill-constrained and Hill-free) from the
results TSVs that plot_curves.py writes to results/.

Requires: pip install python-docx
"""
import math
import re
from pathlib import Path

import pandas as pd
from matplotlib import font_manager
from PIL import ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

# ---------- layout (DXA/twips), portrait A4 with 1" margins ----------
FONT = "Times New Roman"
FONT_PT = 12
TEXT_WIDTH = 11906 - 2 * 1440      # usable page width
CELL_PAD = 2 * 108 + 40            # Word's default left+right cell margins, plus slack

# "_{..}" -> subscript, "^{..}" -> superscript (real Word formatting)
HEADERS = [
    "ID", "IC_{50} (\u00b5M) [95% CI]",
    "Per-expt. IC_{50} (\u00b5M)",
    "pIC_{50}", "Hill", "Top (%)", "Bot. (%)", "R^{2}",
]
DASH = "\u2013"

RESULTS_DIR = Path(__file__).parent / "results"

# Manuscript compound numbers (THQ) for the internal compound IDs
ID_MAP = pd.read_csv(Path(__file__).parent.parent / "compound_id_map.csv").set_index("internal_id")["thq_id"].to_dict()

TABLES = [
    ("ic50_results_hill1.tsv", "IC50_table_hill_constrained.docx", "1", "hill1"),
    ("ic50_results_free.tsv",  "IC50_table_hill_free.docx",        "2", "free"),
]


# ---------- formatting helpers ----------
def sig2(v):
    """Two significant figures, keeping trailing zeros."""
    if v is None or (isinstance(v, float) and math.isnan(v)) or v == "":
        return DASH
    try:
        x = float(v)
    except (TypeError, ValueError):
        return DASH
    if not math.isfinite(x) or x == 0:
        return DASH
    magnitude = math.floor(math.log10(abs(x)))
    decimals = max(0, 1 - magnitude)
    return f"{x:.{decimals}f}"


def num(v, dp):
    if v is None or (isinstance(v, float) and math.isnan(v)) or v == "":
        return DASH
    try:
        return f"{float(v):.{dp}f}"
    except (TypeError, ValueError):
        return DASH


def ci(lo, hi):
    a, b = sig2(lo), sig2(hi)
    return DASH if a == DASH or b == DASH else f"{a}{DASH}{b}"


def ic50_with_ci(val, lo, hi):
    """'0.52 [0.16-1.8]', or just the value when no CI is available."""
    v = sig2(val)
    if v == DASH:
        return DASH
    interval = ci(lo, hi)
    return v if interval == DASH else f"{v} [{interval}]"


def pic50_from(ic50_um):
    """pIC50 derived from the reported IC50 in uM, so the columns agree."""
    try:
        x = float(ic50_um)
    except (TypeError, ValueError):
        return DASH
    if not math.isfinite(x) or x <= 0:
        return DASH
    return f"{6 - math.log10(x):.2f}"


def fold(v):
    if v is None or (isinstance(v, float) and math.isnan(v)) or v == "":
        return DASH
    try:
        return f"\u00d7/\u00f7 {float(v):.2f}"
    except (TypeError, ValueError):
        return DASH


# ---------- low-level cell styling (python-docx has no border API) ----------
def set_cell_border(cell, edge, sz=None, color="000000"):
    """sz in eighths of a point; sz=None removes the border."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    tag = f"w:{edge}"
    el = borders.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        borders.append(el)
    if sz is None:
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
    else:
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), color)


def shade(cell, fill="F2F2F2"):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")       # never "solid" - renders black
    el.set(qn("w:fill"), fill)
    tcPr.append(el)


def add_rich_text(p, text, bold=False):
    """Add text to a paragraph, turning _{..} / ^{..} into sub/superscript runs."""
    for part in re.split(r"([_^]\{[^}]*\})", text):
        if not part:
            continue
        script = part[0] if re.fullmatch(r"[_^]\{[^}]*\}", part) else None
        run = p.add_run(part[2:-1] if script else part)
        run.bold = bold
        run.font.size = Pt(FONT_PT)
        run.font.name = FONT
        run.font.subscript = script == "_"
        run.font.superscript = script == "^"


def no_wrap(cell):
    cell._tc.get_or_add_tcPr().append(OxmlElement("w:noWrap"))


def write_cell(cell, text, width, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.width = Twips(width)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(1)
    pf.space_after = Pt(1)
    add_rich_text(p, text, bold)
    for edge in ("top", "bottom", "left", "right"):
        set_cell_border(cell, edge, None)


def column_widths(rows):
    """Fit each column to its widest data entry, then share the leftover page
    width equally so the table spans the full text width."""
    # Font loaded at FONT_PT*20 px, so measured lengths come out in twips.
    font = ImageFont.truetype(font_manager.findfont(FONT), FONT_PT * 20)
    widths = [max(font.getlength(r[i]) for r in rows) + CELL_PAD
              for i in range(len(HEADERS))]
    spare = TEXT_WIDTH - sum(widths)
    if spare < 0:
        print(f"  warning: data rows are {-spare:.0f} twips wider than the page")
        spare = 0
    return [round(w + spare / len(widths)) for w in widths]


# ---------- main ----------
def build_table(tsv_path, out_path, table_no, hill_mode):
    constrained = hill_mode == "hill1"

    df = pd.read_csv(tsv_path, sep="\t")

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Twips(11906), Twips(16838)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, m, Twips(1440))

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    title = ("Hill slope constrained to 1" if constrained
             else "Hill slope unconstrained")
    h = doc.add_paragraph()
    add_rich_text(
        h,
        f"Supplementary Table {table_no}. Biochemical IC_{{50}} determinations "
        f"against CDK2/Cyclin E1 \u2014 {title}",
        bold=True,
    )
    h.paragraph_format.space_after = Pt(8)

    # body text first, so the column widths can be fitted to it
    rows = []
    for _, r in df.iterrows():
        per_exp = r.get("per_exp_ic50")
        per_exp = DASH if pd.isna(per_exp) or per_exp == "-" else str(per_exp)
        # pIC50 is derived from the REPORTED IC50 so the two always agree;
        # the global-fit pIC50 column would not match the geometric mean.
        rows.append([
            ID_MAP[r["compound"]],
            ic50_with_ci(r["reported_ic50"], r["reported_ci_low"], r["reported_ci_high"]),
            per_exp,
            pic50_from(r["reported_ic50"]),
            "1.00" if constrained else num(r["hill"], 2),
            num(r["top"], 1),
            num(r["bottom"], 1),
            num(r["r2"], 2),
        ])
    cols = column_widths(rows)

    table = doc.add_table(rows=1, cols=len(cols))
    table.autofit = False
    for col, w in zip(table.columns, cols):
        col.width = Twips(w)

    # header
    for i, head in enumerate(HEADERS):
        c = table.rows[0].cells[i]
        write_cell(c, head, cols[i], bold=True)
        shade(c)
        set_cell_border(c, "top", 6)
        set_cell_border(c, "bottom", 4)

    # body
    for ri, vals in enumerate(rows):
        last = ri == len(rows) - 1
        cells = table.add_row().cells
        for i, v in enumerate(vals):
            align = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            write_cell(cells[i], v, cols[i], align=align)
            no_wrap(cells[i])
            set_cell_border(cells[i], "bottom", 6 if last else 2,
                            "000000" if last else "BFBFBF")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)


    doc.save(out_path)
    print(f"wrote {Path(out_path).name}")


def main():
    for tsv, out, table_no, mode in TABLES:
        build_table(RESULTS_DIR / tsv, RESULTS_DIR / out, table_no, mode)


if __name__ == "__main__":
    main()